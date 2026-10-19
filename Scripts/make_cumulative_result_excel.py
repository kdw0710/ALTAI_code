#!/usr/bin/env python3
"""Create and update the cumulative ALTAI chip-test workbook.

The script is intentionally scoped to the project that contains it:

    ALTAI_CHIPTEST_NEW/
      Scripts/make_cumulative_result_excel.py
      testdata/testdata_YYMMDD/...
      Results/ALTAI_chip_results.xlsx

Each execution rebuilds the cumulative ``All chip data``,
``Chip Yield Analysis``, ``Failure Analysis``, and ``NK Chip Details`` sheets,
and creates one static result sheet for every ``testdata_YYMMDD`` batch:

* no DACSCAN_YYMMDD (I_d -4.4), containing only that testdata batch

The top yield table is cumulative.  The lower yield table maps one-to-one to
the batch result sheets.  Manual values entered in the ``Usage`` and
``Current Location`` columns are preserved by corrected physical Chip ID.

Failure analysis is order-independent: for each active no-DACSCAN parameter,
the chip is reclassified with only that parameter excluded.  The workbook
reports which NK chips recover, their recovered grades, and the corresponding
yield change for the cumulative data and for every testdata batch.
"""

from __future__ import annotations

import argparse
import copy
import json
import math
import os
import re
from collections import Counter
from itertools import groupby
from pathlib import Path, PurePosixPath
from typing import Any, Callable
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter


SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
TESTDATA_ROOT = PROJECT_ROOT / "testdata"
RESULTS_ROOT = PROJECT_ROOT / "Results"

DEFAULT_OUTPUT_NAME = "ALTAI_chip_results.xlsx"
GRADE_ORDER = ("GOLD", "SILVER", "BRONZE", "BRONZENOVBB")
YIELD_GRADE_ORDER = (*GRADE_ORDER, "NK", "CONTACT_FAIL")
YIELD_MODES = (
    ("no DACSCAN", "no_dac_grade"),
    ("Only Power", "power_grade"),
)
TEST_GROUP_LABELS = {
    "PowerTest": "Powering Test",
    "BackbiasIV": "Back-bias I-V Curve",
    "ChipID": "Chip ID Test",
    "Registers": "Register Test",
    "FIFOTest": "Memory Test",
    "SEU": "SEU Test",
    "DACScan": "DAC Scan",
    "DataPort": "Data Port Test",
    "AnalogueScanDataPort": "Analogue Scan",
    "DigitalScan": "Digital Scan",
    "ThresholdScanDataPortQueue": "Threshold Scan",
    "FakeHitRate": "Fake Hit Rate",
}
TEST_GROUP_ORDER = tuple(TEST_GROUP_LABELS)
TEST_GROUP_INDEX = {
    group: index
    for index, group in enumerate(TEST_GROUP_ORDER)
}
NO_DAC_EXCLUDED_GROUPS = {"DACScan", "BackbiasIV"}

# This correction is deliberately restricted to this one imported data batch.
# The folder names are never changed; only the Chip ID written to the workbook
# is corrected.
BATCH_ID_CORRECTIONS = {
    "testdata_260318": (
        (1772722654, 1772732321, "0RUZS07B_", "0RUZS07T_"),
    ),
}

# The first part of the April delivery keeps the ATE export hierarchy and
# writes the tray position directly after the lot ID (for example,
# ``51PK417BA1``).  Limit this normalization to that one imported batch so an
# unrelated future chip ID is never changed implicitly.
BATCH_ID_REGEX_CORRECTIONS = {
    "testdata_260416": (
        (re.compile(r"^(51PK417B)([ABC][1-8])$"), r"\1_\2"),
    ),
}

RUN_FOLDER_PATTERN = re.compile(r"^\d{8}-\d{6}$")

TITLE_FILL = PatternFill("solid", fgColor="1F4E78")
SUBTITLE_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FILL = PatternFill("solid", fgColor="5B9BD5")
PARAM_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
MANUAL_HEADER_FILL = PatternFill("solid", fgColor="F4B183")
MANUAL_CELL_FILL = PatternFill("solid", fgColor="FFF2CC")
EXCLUDED_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
GROUP_HEADER_FILL = PatternFill("solid", fgColor="2F75B5")
EXCLUDED_GROUP_FILL = PatternFill("solid", fgColor="7F8C8D")
THIN_GRAY = Side(style="thin", color="D9E2F3")
WHITE_SEPARATOR = Side(style="thin", color="FFFFFF")
TABLE_BORDER = Border(bottom=THIN_GRAY)
MANUAL_CELL_BORDER = Border(
    left=THIN_GRAY,
    right=THIN_GRAY,
    bottom=THIN_GRAY,
)

GRADE_FILL = {
    "GOLD": PatternFill("solid", fgColor="FFD966"),
    "SILVER": PatternFill("solid", fgColor="D9E1F2"),
    "BRONZE": PatternFill("solid", fgColor="C68642"),
    "BRONZENOVBB": PatternFill("solid", fgColor="F8CBAD"),
    "NK": PatternFill("solid", fgColor="F4CCCC"),
    "CONTACT_FAIL": PatternFill("solid", fgColor="A6A6A6"),
    "N/A": PatternFill("solid", fgColor="E7E6E6"),
    "MISSING": PatternFill("solid", fgColor="E7E6E6"),
}

GRADE_FONT = {
    "GOLD": "7F6000",
    "SILVER": "404040",
    "BRONZE": "4F2B0F",
    "BRONZENOVBB": "843C0C",
    "NK": "9C0006",
    "CONTACT_FAIL": "FFFFFF",
    "N/A": "7F7F7F",
    "MISSING": "7F7F7F",
}

FAILURE_FILL = PatternFill("solid", fgColor="FCE4D6")
RECOVERY_FILL = PatternFill("solid", fgColor="E2F0D9")
INFO_FILL = PatternFill("solid", fgColor="EDEDED")
GRADE_SHORT_LABELS = {
    "GOLD": "G",
    "SILVER": "S",
    "BRONZE": "B",
    "BRONZENOVBB": "BNV",
}


def validate_project_layout() -> None:
    """Fail closed if the script is copied away from the intended project."""
    if PROJECT_ROOT.name != "ALTAI_CHIPTEST_NEW":
        raise RuntimeError(
            "This script must be run from ALTAI_CHIPTEST_NEW/Scripts. "
            f"Detected project root: {PROJECT_ROOT}"
        )
    if not TESTDATA_ROOT.is_dir():
        raise RuntimeError(f"Project testdata directory not found: {TESTDATA_ROOT}")
    if not (SCRIPT_DIR / "classification").is_dir():
        raise RuntimeError(
            f"Classification criteria directory not found: "
            f"{SCRIPT_DIR / 'classification'}"
        )


def get_output_path(output_name: str) -> Path:
    candidate = Path(output_name)
    if candidate.name != output_name or candidate.suffix.lower() != ".xlsx":
        raise ValueError("--output-name must be a simple .xlsx file name.")
    output = (RESULTS_ROOT / candidate.name).resolve()
    results_root = RESULTS_ROOT.resolve()
    if output.parent != results_root:
        raise ValueError(f"Output must remain inside {results_root}")
    return output


def load_criteria() -> tuple[dict[str, dict[str, list[float]]], int | None]:
    cuts: dict[str, dict[str, list[float]]] = {}
    for grade in GRADE_ORDER:
        cut_path = SCRIPT_DIR / "classification" / f"{grade}.json"
        with cut_path.open("r", encoding="utf-8") as handle:
            cuts[grade] = json.load(handle)

    version = None
    init_path = SCRIPT_DIR / "classification" / "__init__.py"
    match = re.search(
        r"^\s*VERSION\s*=\s*(\d+)\s*$",
        init_path.read_text(encoding="utf-8"),
        flags=re.MULTILINE,
    )
    if match:
        version = int(match.group(1))
    return cuts, version


def flatten(value: Any, prefix: str = "", output: dict[str, Any] | None = None):
    if output is None:
        output = {}
    if isinstance(value, dict):
        for key, child in value.items():
            next_prefix = f"{prefix}.{key}" if prefix else key
            flatten(child, next_prefix, output)
    else:
        output[prefix] = value
    return output


def normalize_value(value: Any) -> Any:
    """Keep spreadsheet values valid while retaining missing-data meaning."""
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def correct_chip_id(
    chip_id: str,
    outer_folder: str,
    batch_name: str,
) -> tuple[str, bool]:
    corrections = BATCH_ID_CORRECTIONS.get(batch_name, ())
    try:
        outer_number = int(outer_folder)
    except ValueError:
        return chip_id, False

    corrected = chip_id
    for start, end, old_prefix, new_prefix in corrections:
        if start <= outer_number <= end and corrected.startswith(old_prefix):
            corrected = corrected.replace(old_prefix, new_prefix, 1)

    for pattern, replacement in BATCH_ID_REGEX_CORRECTIONS.get(
        batch_name,
        (),
    ):
        corrected = pattern.sub(replacement, corrected)
    return corrected, corrected != chip_id


def discover_run_folders() -> list[tuple[Path, Path, Path, Path]]:
    """Return (batch, outer, chip, run) directories in deterministic order.

    Two source layouts are supported without moving or renaming raw data:

    * the normal ``batch/outer/chip/run`` layout used by the existing batches;
    * the ATE export layout
      ``batch/session/CHDO/Chips/lot/chip/CERN_PC/run`` used by part 1 of the
      April 2026 delivery.
    """
    discovered: list[tuple[Path, Path, Path, Path]] = []
    seen_runs: set[Path] = set()

    def add_run(batch: Path, outer: Path, chip: Path, run: Path) -> None:
        resolved = run.resolve()
        if resolved in seen_runs:
            return
        seen_runs.add(resolved)
        discovered.append((batch, outer, chip, run))

    for batch in sorted(TESTDATA_ROOT.glob("testdata_*")):
        if not batch.is_dir():
            continue
        for outer in sorted(path for path in batch.iterdir() if path.is_dir()):
            # Standard project layout.  This remains the primary path and is
            # intentionally unchanged for all existing testdata batches.
            for chip in sorted(path for path in outer.iterdir() if path.is_dir()):
                for run in sorted(path for path in chip.iterdir() if path.is_dir()):
                    if RUN_FOLDER_PATTERN.fullmatch(run.name):
                        add_run(batch, outer, chip, run)

            # ATE export layout found in ``ALTAI_test_part_1`` of the April
            # delivery.  The session directory is retained as ``outer`` and
            # the actual chip directory below ``Chips/<lot>`` supplies the ID.
            chips_root = outer / "CHDO" / "Chips"
            if not chips_root.is_dir():
                continue
            for lot in sorted(
                path for path in chips_root.iterdir() if path.is_dir()
            ):
                for chip in sorted(
                    path for path in lot.iterdir() if path.is_dir()
                ):
                    cern_pc = chip / "CERN_PC"
                    if not cern_pc.is_dir():
                        continue
                    for run in sorted(
                        path for path in cern_pc.iterdir() if path.is_dir()
                    ):
                        if RUN_FOLDER_PATTERN.fullmatch(run.name):
                            add_run(batch, outer, chip, run)
    return discovered


def collect_records() -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for batch, outer, chip, run in discover_run_folders():
        original_chip_id = chip.name
        chip_id, corrected = correct_chip_id(
            original_chip_id,
            outer.name,
            batch.name,
        )
        analyse_path = run / "analyse_output.json"
        contact_fail_path = run / "CONTACT_FAIL"

        parameters: dict[str, Any] = {}
        if analyse_path.is_file():
            try:
                with analyse_path.open("r", encoding="utf-8") as handle:
                    parameters = flatten(json.load(handle))
                status = "ANALYZED"
            except (OSError, ValueError, TypeError) as exc:
                print(f"[WARN] Cannot read {analyse_path}: {exc}")
                status = "INVALID_ANALYSE_OUTPUT"
        elif contact_fail_path.is_file():
            status = "CONTACT_FAIL"
        else:
            status = "NO_ANALYSE_OUTPUT"

        records.append(
            {
                "chip_id": chip_id,
                "original_chip_id": original_chip_id if corrected else "",
                "batch": batch.name,
                "wafer": outer.name,
                "test_date": run.name,
                "status": status,
                "params": parameters,
                "source": str(run.relative_to(TESTDATA_ROOT)),
            }
        )

    records.sort(
        key=lambda record: (
            record["chip_id"],
            record["test_date"],
            record["batch"],
            record["wafer"],
        )
    )
    return records


def latest_per_chip(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    for record in records:
        chip_id = record["chip_id"]
        current = latest.get(chip_id)
        selection_key = (
            record["test_date"],
            record["batch"],
            record["wafer"],
        )
        if current is None:
            latest[chip_id] = record
            continue
        current_key = (
            current["test_date"],
            current["batch"],
            current["wafer"],
        )
        if selection_key > current_key:
            latest[chip_id] = record

    return sorted(
        latest.values(),
        key=lambda record: (
            record["batch"],
            record["wafer"],
            record["chip_id"],
        ),
    )


def exclude_dacscan_and_backbias(parameter: str) -> bool:
    return not (
        parameter.startswith("DACScan.")
        or parameter.startswith("BackbiasIV.")
    )


def include_power_only(parameter: str) -> bool:
    return parameter.startswith("PowerTest.")


def grade_chip(
    parameters: dict[str, Any],
    cuts: dict[str, dict[str, list[float]]],
    parameter_filter: Callable[[str], bool],
) -> str:
    for grade in GRADE_ORDER:
        passed = True
        for parameter, (low, high) in cuts[grade].items():
            if not parameter_filter(parameter):
                continue
            value = parameters.get(parameter)
            if (
                value is None
                or not isinstance(value, (int, float))
                or isinstance(value, bool)
                or math.isnan(value)
                or value < low
                or value > high
            ):
                passed = False
                break
        if passed:
            return grade
    return "NK"


def grade_value(
    parameter: str,
    value: Any,
    cuts: dict[str, dict[str, list[float]]],
) -> str:
    if (
        value is None
        or not isinstance(value, (int, float))
        or isinstance(value, bool)
        or (isinstance(value, float) and math.isnan(value))
    ):
        return "MISSING"

    has_cut = False
    for grade in GRADE_ORDER:
        cut = cuts[grade].get(parameter)
        if cut is None:
            continue
        has_cut = True
        if cut[0] <= value <= cut[1]:
            return grade
    return "NK" if has_cut else "N/A"


def is_numeric_measurement(value: Any) -> bool:
    return (
        isinstance(value, (int, float))
        and not isinstance(value, bool)
        and not (isinstance(value, float) and math.isnan(value))
    )


def no_dac_active_parameters(
    cuts: dict[str, dict[str, list[float]]],
) -> list[str]:
    return [
        parameter
        for parameter in all_parameters(cuts)
        if exclude_dacscan_and_backbias(parameter)
    ]


def grade_failure_details(
    parameters: dict[str, Any],
    cuts: dict[str, dict[str, list[float]]],
) -> tuple[dict[str, dict[str, Any]], dict[str, str]]:
    """Return failed grade cuts grouped by parameter and one state per parameter."""
    failed: dict[str, dict[str, Any]] = {}
    states: dict[str, str] = {}
    for grade in GRADE_ORDER:
        for parameter, (low, high) in cuts[grade].items():
            if not exclude_dacscan_and_backbias(parameter):
                continue
            value = parameters.get(parameter)
            if not is_numeric_measurement(value):
                state = "MISSING"
            elif value < low or value > high:
                state = "OUT_OF_RANGE"
            else:
                continue

            entry = failed.setdefault(
                parameter,
                {
                    "grades": [],
                    "states": set(),
                },
            )
            entry["grades"].append(grade)
            entry["states"].add(state)

    for parameter, detail in failed.items():
        states[parameter] = (
            "MISSING"
            if "MISSING" in detail["states"]
            else "OUT_OF_RANGE"
        )
    return failed, states


def grade_without_parameter(
    parameters: dict[str, Any],
    cuts: dict[str, dict[str, list[float]]],
    excluded_parameter: str,
) -> str:
    return grade_chip(
        parameters,
        cuts,
        lambda parameter: (
            exclude_dacscan_and_backbias(parameter)
            and parameter != excluded_parameter
        ),
    )


def add_failure_details(
    record: dict[str, Any],
    cuts: dict[str, dict[str, list[float]]],
) -> None:
    record["nk_reason_count"] = 0
    record["nk_failed_criteria"] = ""
    record["nk_cause_summary"] = ""
    record["missing_parameters"] = ""
    record["missing_parameter_count"] = 0
    record["recovered_if_excluded"] = ""
    record["nk_failure_states"] = {}
    record["recovery_by_parameter"] = {}

    if record["status"] == "CONTACT_FAIL":
        return
    if record["no_dac_grade"] != "NK":
        return
    if record["status"] != "ANALYZED":
        record["nk_reason_count"] = 1
        record["nk_failed_criteria"] = record["status"]
        record["nk_cause_summary"] = record["status"]
        record["missing_parameters"] = "No valid analyse_output.json"
        record["missing_parameter_count"] = 1
        return

    failed, states = grade_failure_details(
        record["params"],
        cuts,
    )
    record["nk_failure_states"] = states
    record["nk_reason_count"] = len(failed)
    failed_labels = []
    for parameter in no_dac_active_parameters(cuts):
        if parameter not in failed:
            continue
        grade_labels = ",".join(
            GRADE_SHORT_LABELS[grade]
            for grade in failed[parameter]["grades"]
        )
        state = states[parameter].replace("_", " ")
        failed_labels.append(
            f"{parameter} [{grade_labels}; {state}]"
        )
    record["nk_failed_criteria"] = "; ".join(failed_labels)
    record["missing_parameters"] = "; ".join(
        parameter
        for parameter in no_dac_active_parameters(cuts)
        if states.get(parameter) == "MISSING"
    )
    missing_count = sum(state == "MISSING" for state in states.values())
    out_of_range_count = sum(
        state == "OUT_OF_RANGE"
        for state in states.values()
    )
    record["missing_parameter_count"] = missing_count
    summary_parts = []
    if missing_count:
        summary_parts.append(f"Missing {missing_count}")
    if out_of_range_count:
        summary_parts.append(f"Out of range {out_of_range_count}")
    record["nk_cause_summary"] = "; ".join(summary_parts)

    recoveries: dict[str, str] = {}
    for parameter in no_dac_active_parameters(cuts):
        recovered_grade = grade_without_parameter(
            record["params"],
            cuts,
            parameter,
        )
        if recovered_grade in GRADE_ORDER:
            recoveries[parameter] = recovered_grade
    record["recovery_by_parameter"] = recoveries
    record["recovered_if_excluded"] = "; ".join(
        f"{parameter} → {grade}"
        for parameter, grade in recoveries.items()
    )


def add_grades(
    records: list[dict[str, Any]],
    cuts: dict[str, dict[str, list[float]]],
) -> list[dict[str, Any]]:
    output = []
    for record in records:
        updated = copy.copy(record)
        if record["status"] == "CONTACT_FAIL":
            updated["no_dac_grade"] = "CONTACT_FAIL"
            updated["power_grade"] = "CONTACT_FAIL"
        else:
            updated["no_dac_grade"] = grade_chip(
                record["params"],
                cuts,
                exclude_dacscan_and_backbias,
            )
            updated["power_grade"] = grade_chip(
                record["params"],
                cuts,
                include_power_only,
            )
        add_failure_details(updated, cuts)
        output.append(updated)
    return output


def build_batch_records(
    records: list[dict[str, Any]],
    cuts: dict[str, dict[str, list[float]]],
) -> dict[str, list[dict[str, Any]]]:
    """Return one independently deduplicated, graded record set per batch."""
    batches: dict[str, list[dict[str, Any]]] = {}
    batch_names = sorted({record["batch"] for record in records})
    for batch_name in batch_names:
        batch_source = [
            record
            for record in records
            if record["batch"] == batch_name
        ]
        batches[batch_name] = add_grades(
            latest_per_chip(batch_source),
            cuts,
        )
    return batches


def batch_result_sheet_name(batch_name: str) -> str:
    prefix = "testdata_"
    if not batch_name.startswith(prefix):
        raise ValueError(f"Unexpected testdata batch name: {batch_name}")
    batch_date = batch_name[len(prefix):]
    if not re.fullmatch(r"\d{6}", batch_date):
        raise ValueError(f"Unexpected testdata batch date: {batch_name}")
    sheet_name = f"no DACSCAN_{batch_date} (I_d -4.4)"
    if len(sheet_name) > 31:
        raise RuntimeError(
            f"Generated sheet name exceeds Excel's 31-character limit: "
            f"{sheet_name}"
        )
    return sheet_name


def all_parameters(cuts: dict[str, dict[str, list[float]]]) -> list[str]:
    parameters = {
        parameter
        for grade_cuts in cuts.values()
        for parameter in grade_cuts
    }
    return sorted(
        parameters,
        key=lambda parameter: (
            TEST_GROUP_INDEX.get(parameter_group(parameter), 999),
            parameter,
        ),
    )


def parameter_group(parameter: str) -> str:
    return parameter.split(".", 1)[0]


def extract_manual_fields(workbook) -> dict[str, dict[str, str]]:
    if "All chip data" not in workbook.sheetnames:
        return {}
    sheet = workbook["All chip data"]
    header_row = None
    headers: dict[Any, int] = {}
    for row in range(1, min(sheet.max_row, 10) + 1):
        candidate = {
            sheet.cell(row, column).value: column
            for column in range(1, sheet.max_column + 1)
            if sheet.cell(row, column).value not in (None, "")
        }
        if "Original Chip ID" in candidate or "Chip ID" in candidate:
            if (
                "Usage" in candidate
                or "Current Location" in candidate
                or "Usage / Current Location" in candidate
            ):
                header_row = row
                headers = candidate
                break
    if header_row is None:
        return {}

    chip_column = headers.get("Original Chip ID") or headers.get("Chip ID")
    usage_column = headers.get("Usage")
    location_column = headers.get("Current Location")
    legacy_combined_column = headers.get("Usage / Current Location")
    if not chip_column or not (
        usage_column or location_column or legacy_combined_column
    ):
        return {}

    values: dict[str, dict[str, str]] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        chip_id = sheet.cell(row, chip_column).value
        if not chip_id:
            continue
        usage = sheet.cell(row, usage_column).value if usage_column else None
        location = (
            sheet.cell(row, location_column).value if location_column else None
        )
        if legacy_combined_column and usage in (None, ""):
            usage = sheet.cell(row, legacy_combined_column).value
        fields = {
            "usage": "" if usage in (None, "") else str(usage),
            "location": "" if location in (None, "") else str(location),
        }
        if fields["usage"] or fields["location"]:
            values[str(chip_id)] = fields
    return values


def _zip_cell_value(
    cell,
    shared_strings: list[str],
    namespace: dict[str, str],
) -> Any:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(
            text.text or ""
            for text in cell.iterfind(".//m:t", namespace)
        )
    value_element = cell.find("m:v", namespace)
    if value_element is None:
        return None
    value = value_element.text or ""
    if cell_type == "s":
        return shared_strings[int(value)]
    if cell_type == "b":
        return value == "1"
    return value


def load_usage_source(path: Path | None) -> dict[str, dict[str, str]]:
    """Read the small manual table without loading the workbook calculation chain."""
    if path is None:
        return {}
    source = path.expanduser().resolve()
    if not source.is_file() or source.suffix.lower() != ".xlsx":
        raise ValueError(f"Usage source must be an existing .xlsx file: {source}")

    main_namespace = {
        "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "r": (
            "http://schemas.openxmlformats.org/officeDocument/2006/"
            "relationships"
        ),
    }
    relationship_namespace = {
        "r": "http://schemas.openxmlformats.org/package/2006/relationships"
    }

    with ZipFile(source) as archive:
        shared_strings: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared_root = ElementTree.fromstring(
                archive.read("xl/sharedStrings.xml")
            )
            for item in shared_root.findall("m:si", main_namespace):
                shared_strings.append(
                    "".join(
                        text.text or ""
                        for text in item.iterfind(".//m:t", main_namespace)
                    )
                )

        workbook_root = ElementTree.fromstring(
            archive.read("xl/workbook.xml")
        )
        sheet_relationship_id = None
        for sheet in workbook_root.findall(
            ".//m:sheets/m:sheet",
            main_namespace,
        ):
            if sheet.attrib.get("name") == "Chip Usage":
                sheet_relationship_id = sheet.attrib.get(
                    f"{{{main_namespace['r']}}}id"
                )
                break
        if not sheet_relationship_id:
            raise ValueError(
                f"'Chip Usage' sheet not found in usage source: {source}"
            )

        relationships_root = ElementTree.fromstring(
            archive.read("xl/_rels/workbook.xml.rels")
        )
        sheet_target = None
        for relationship in relationships_root.findall(
            "r:Relationship",
            relationship_namespace,
        ):
            if relationship.attrib.get("Id") == sheet_relationship_id:
                sheet_target = relationship.attrib.get("Target")
                break
        if not sheet_target:
            raise ValueError(
                f"Cannot resolve 'Chip Usage' sheet in usage source: {source}"
            )

        target_path = PurePosixPath(sheet_target.lstrip("/"))
        if not str(target_path).startswith("xl/"):
            target_path = PurePosixPath("xl") / target_path
        sheet_root = ElementTree.fromstring(archive.read(str(target_path)))

        table_rows: list[dict[int, Any]] = []
        for row in sheet_root.findall(
            ".//m:sheetData/m:row",
            main_namespace,
        ):
            values: dict[int, Any] = {}
            for cell in row.findall("m:c", main_namespace):
                match = re.match(r"([A-Z]+)", cell.attrib.get("r", ""))
                if not match:
                    continue
                column = column_index_from_string(match.group(1))
                values[column] = _zip_cell_value(
                    cell,
                    shared_strings,
                    main_namespace,
                )
            if values:
                table_rows.append(values)

    required_headers = {
        "Final Usage",
        "Current Location",
        "Chip ID",
    }
    header_map: dict[str, int] | None = None
    header_index = None
    for index, row in enumerate(table_rows):
        candidate = {
            str(value): column
            for column, value in row.items()
            if value not in (None, "")
        }
        if required_headers.issubset(candidate):
            header_map = candidate
            header_index = index
            break
    if header_map is None or header_index is None:
        raise ValueError(
            "Usage source requires Final Usage, Current Location, and Chip ID "
            f"headers: {source}"
        )

    values: dict[str, dict[str, str]] = {}
    for row in table_rows[header_index + 1 :]:
        chip_id = row.get(header_map["Chip ID"])
        if chip_id in (None, ""):
            continue
        key = str(chip_id).strip()
        current = values.setdefault(key, {"usage": "", "location": ""})
        usage = row.get(header_map["Final Usage"])
        location = row.get(header_map["Current Location"])
        if usage not in (None, ""):
            current["usage"] = str(usage)
        if location not in (None, ""):
            current["location"] = str(location)
    return values


def merge_manual_fields(
    imported: dict[str, dict[str, str]],
    existing: dict[str, dict[str, str]],
) -> dict[str, dict[str, str]]:
    """Existing workbook edits take precedence field by field."""
    merged = {
        chip_id: {
            "usage": fields.get("usage", ""),
            "location": fields.get("location", ""),
        }
        for chip_id, fields in imported.items()
    }
    for chip_id, fields in existing.items():
        target = merged.setdefault(
            chip_id,
            {"usage": "", "location": ""},
        )
        if fields.get("usage"):
            target["usage"] = fields["usage"]
        if fields.get("location"):
            target["location"] = fields["location"]
    return merged


def set_title(sheet, title: str, subtitle: str, last_column: int) -> None:
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(f"A1:{last_letter}1")
    sheet["A1"] = title
    sheet["A1"].fill = TITLE_FILL
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=15)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 26

    sheet.merge_cells(f"A2:{last_letter}2")
    sheet["A2"] = subtitle
    sheet["A2"].fill = SUBTITLE_FILL
    sheet["A2"].font = Font(color="1F1F1F", italic=True, size=10)
    sheet["A2"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[2].height = 34


def style_header(
    cell,
    manual: bool = False,
    parameter: bool = False,
    excluded: bool = False,
) -> None:
    if manual:
        cell.fill = MANUAL_HEADER_FILL
    elif excluded:
        cell.fill = EXCLUDED_HEADER_FILL
    elif parameter:
        cell.fill = PARAM_HEADER_FILL
    else:
        cell.fill = HEADER_FILL
    cell.font = Font(
        color="000000" if (manual or parameter or excluded) else "FFFFFF",
        bold=True,
        size=10,
    )
    cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )
    cell.border = TABLE_BORDER


def write_group_header(
    sheet,
    start_column: int,
    end_column: int,
    label: str,
    excluded: bool = False,
) -> None:
    start_letter = get_column_letter(start_column)
    end_letter = get_column_letter(end_column)
    if end_column > start_column:
        sheet.merge_cells(f"{start_letter}3:{end_letter}3")
    anchor = sheet.cell(3, start_column)
    anchor.value = label
    fill = EXCLUDED_GROUP_FILL if excluded else GROUP_HEADER_FILL
    for column in range(start_column, end_column + 1):
        cell = sheet.cell(3, column)
        cell.fill = fill
        cell.border = Border(
            left=WHITE_SEPARATOR if column == start_column else Side(),
            right=WHITE_SEPARATOR if column == end_column else Side(),
            bottom=WHITE_SEPARATOR,
        )
    anchor.font = Font(color="FFFFFF", bold=True, size=11)
    anchor.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


def write_parameter_group_headers(
    sheet,
    parameters: list[str],
    parameter_start: int,
    excluded_groups: set[str] | None = None,
) -> None:
    excluded_groups = excluded_groups or set()
    indexed = list(enumerate(parameters, parameter_start))
    for group, items_iterator in groupby(
        indexed,
        key=lambda item: parameter_group(item[1]),
    ):
        items = list(items_iterator)
        start_column = items[0][0]
        end_column = items[-1][0]
        label = TEST_GROUP_LABELS.get(group, group)
        excluded = group in excluded_groups
        if excluded:
            label = f"{label}\n(shown; excluded from Global)"
        write_group_header(
            sheet,
            start_column,
            end_column,
            label,
            excluded=excluded,
        )


def style_grade_cell(cell, grade: str) -> None:
    cell.fill = GRADE_FILL.get(grade, GRADE_FILL["N/A"])
    cell.font = Font(
        color=GRADE_FONT.get(grade, "000000"),
        bold=True,
    )
    cell.alignment = Alignment(horizontal="center", vertical="center")


def configure_sheet(
    sheet,
    last_column: int,
    freeze_panes: str,
    widths: dict[int, float],
    header_row: int = 3,
) -> None:
    sheet.freeze_panes = freeze_panes
    sheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(last_column)}{sheet.max_row}"
    )
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[header_row].height = 54
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width


def write_all_chip_data(
    workbook,
    records: list[dict[str, Any]],
    cuts: dict[str, dict[str, list[float]]],
    criteria_version: int | None,
    manual_fields: dict[str, dict[str, str]],
) -> None:
    if "All chip data" in workbook.sheetnames:
        del workbook["All chip data"]
    sheet = workbook.create_sheet("All chip data", 0)

    parameters = all_parameters(cuts)
    headers = [
        "Original Chip ID",
        "Usage",
        "Current Location",
        "Wrong Written Chip ID",
        "Data Batch",
        "Wafer / Tray",
        "Test Date",
        "Status",
        "no DACSCAN Grade",
        "Only Power Grade",
        "NK Reason Count",
        "NK Cause Summary",
        "Missing Parameter Count",
        "Recovered if Excluded",
        *parameters,
        "Source Folder",
    ]
    set_title(
        sheet,
        "ALTAI Chip Test — Cumulative Chip Data",
        (
            "Yellow columns are for manual Usage and Current Location entries. "
            "Entries are preserved by Original Chip ID on every rerun. "
            f"Criteria version: {criteria_version if criteria_version is not None else 'unknown'}; "
            "DACScan and Backbias are excluded from the no-DACSCAN grade. "
            "NK cause columns show failed grade cuts and the order-independent "
            "grade recovered by excluding one active parameter. "
            "Cell colors show each measured parameter's grade."
        ),
        len(headers),
    )

    write_group_header(sheet, 1, 4, "Chip Identification & Tracking")
    write_group_header(sheet, 5, 8, "Test Metadata")
    write_group_header(sheet, 9, 14, "Final Classification & NK Cause")
    parameter_start = 15
    write_parameter_group_headers(
        sheet,
        parameters,
        parameter_start,
    )
    source_column = parameter_start + len(parameters)
    write_group_header(sheet, source_column, source_column, "Source")
    sheet.row_dimensions[3].height = 34

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(4, column, header)
        style_header(
            cell,
            manual=header in {"Usage", "Current Location"},
            parameter=header in parameters,
        )

    for row_number, record in enumerate(records, 5):
        manual = manual_fields.get(
            record["chip_id"],
            {"usage": "", "location": ""},
        )
        metadata = [
            record["chip_id"],
            manual["usage"],
            manual["location"],
            record["original_chip_id"],
            record["batch"],
            record["wafer"],
            record["test_date"],
            record["status"],
            record["no_dac_grade"],
            record["power_grade"],
            record["nk_reason_count"] if record["no_dac_grade"] == "NK" else "",
            record["nk_cause_summary"],
            (
                record["missing_parameter_count"]
                if record["no_dac_grade"] == "NK"
                else ""
            ),
            record["recovered_if_excluded"],
        ]
        for column, value in enumerate(metadata, 1):
            cell = sheet.cell(row_number, column, normalize_value(value))
            cell.border = (
                MANUAL_CELL_BORDER
                if column in {2, 3}
                else TABLE_BORDER
            )
            cell.alignment = Alignment(
                horizontal="left" if column <= 4 or column in {12, 14} else "center",
                vertical="center",
                wrap_text=column in {2, 3, 12, 14},
            )

        sheet.cell(row_number, 2).fill = MANUAL_CELL_FILL
        sheet.cell(row_number, 3).fill = MANUAL_CELL_FILL
        if len(manual["usage"]) > 28 or len(manual["location"]) > 20:
            sheet.row_dimensions[row_number].height = 30
        style_grade_cell(sheet.cell(row_number, 9), record["no_dac_grade"])
        style_grade_cell(sheet.cell(row_number, 10), record["power_grade"])
        if record["no_dac_grade"] == "NK":
            sheet.cell(row_number, 11).fill = FAILURE_FILL
            sheet.cell(row_number, 12).fill = FAILURE_FILL
            sheet.cell(row_number, 13).fill = INFO_FILL
            if record["recovered_if_excluded"]:
                sheet.cell(row_number, 14).fill = RECOVERY_FILL
                sheet.cell(row_number, 14).font = Font(
                    color="375623",
                    bold=True,
                )

        for offset, parameter in enumerate(parameters):
            raw_value = record["params"].get(parameter)
            value = normalize_value(raw_value)
            cell = sheet.cell(row_number, parameter_start + offset, value)
            cell.border = TABLE_BORDER
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "0.0000"
                cell.alignment = Alignment(horizontal="right")
            parameter_grade = grade_value(parameter, raw_value, cuts)
            if parameter_grade not in {"N/A", "MISSING"}:
                cell.fill = GRADE_FILL[parameter_grade]
                cell.font = Font(color=GRADE_FONT[parameter_grade])

        source_cell = sheet.cell(row_number, source_column, record["source"])
        source_cell.border = TABLE_BORDER
        source_cell.font = Font(color="7F7F7F", size=9)

    widths = {
        1: 18,
        2: 34,
        3: 24,
        4: 20,
        5: 20,
        6: 18,
        7: 20,
        8: 18,
        9: 19,
        10: 19,
        11: 16,
        12: 28,
        13: 20,
        14: 42,
    }
    for column in range(parameter_start, parameter_start + len(parameters)):
        widths[column] = 20
    widths[parameter_start + len(parameters)] = 50
    configure_sheet(
        sheet,
        len(headers),
        "E5",
        widths,
        header_row=4,
    )
    sheet.sheet_properties.tabColor = "4472C4"


def summary_text(
    records: list[dict[str, Any]],
    grade_key: str,
    criteria_version: int | None,
) -> str:
    status_counts = Counter(record["status"] for record in records)
    grade_counts = Counter(record[grade_key] for record in records)
    return (
        f"Total {len(records)} | ANALYZED {status_counts.get('ANALYZED', 0)} | "
        f"CONTACT_FAIL {status_counts.get('CONTACT_FAIL', 0)} | "
        f"GOLD {grade_counts.get('GOLD', 0)} | "
        f"SILVER {grade_counts.get('SILVER', 0)} | "
        f"BRONZE {grade_counts.get('BRONZE', 0)} | "
        f"BRONZENOVBB {grade_counts.get('BRONZENOVBB', 0)} | "
        f"NK {grade_counts.get('NK', 0)} | "
        f"CONTACT_FAIL {grade_counts.get('CONTACT_FAIL', 0)} | "
        f"Criteria version {criteria_version if criteria_version is not None else 'unknown'}"
    )


def write_snapshot(
    workbook,
    sheet_name: str,
    index: int,
    title: str,
    records: list[dict[str, Any]],
    grade_key: str,
    parameters: list[str],
    cuts: dict[str, dict[str, list[float]]],
    criteria_version: int | None,
    excluded_groups: set[str] | None = None,
) -> None:
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    sheet = workbook.create_sheet(sheet_name, index)

    headers = [
        "Original Chip ID",
        "Wrong Written Chip ID",
        "Data Batch",
        "Wafer / Tray",
        "Test Date",
        "Status",
        "Global",
        "NK Reason Count",
        "NK Cause Summary",
        "Missing Parameter Count",
        "Recovered if Excluded",
        *parameters,
    ]
    excluded_groups = excluded_groups or set()
    subtitle = summary_text(records, grade_key, criteria_version)
    if excluded_groups:
        subtitle += (
            " | DAC Scan and Back-bias I-V Curve values are shown for "
            "reference but excluded from the Global grade."
        )
    set_title(
        sheet,
        title,
        subtitle,
        len(headers),
    )

    write_group_header(sheet, 1, 2, "Chip Identification")
    write_group_header(sheet, 3, 6, "Test Metadata")
    write_group_header(sheet, 7, 11, "Final Classification & NK Cause")
    parameter_start = 12
    write_parameter_group_headers(
        sheet,
        parameters,
        parameter_start,
        excluded_groups=excluded_groups,
    )
    sheet.row_dimensions[3].height = 42

    for column, header in enumerate(headers, 1):
        cell = sheet.cell(4, column, header)
        group = parameter_group(header) if header in parameters else ""
        style_header(
            cell,
            parameter=header in parameters,
            excluded=group in excluded_groups,
        )

    for row_number, record in enumerate(records, 5):
        metadata = [
            record["chip_id"],
            record["original_chip_id"],
            record["batch"],
            record["wafer"],
            record["test_date"],
            record["status"],
            record[grade_key],
            record["nk_reason_count"] if record[grade_key] == "NK" else "",
            record["nk_cause_summary"],
            (
                record["missing_parameter_count"]
                if record[grade_key] == "NK"
                else ""
            ),
            record["recovered_if_excluded"],
        ]
        for column, value in enumerate(metadata, 1):
            cell = sheet.cell(row_number, column, normalize_value(value))
            cell.border = TABLE_BORDER
            cell.alignment = Alignment(
                horizontal="left" if column <= 2 or column in {9, 11} else "center",
                vertical="center",
                wrap_text=column in {9, 11},
            )
        style_grade_cell(sheet.cell(row_number, 7), record[grade_key])
        if record[grade_key] == "NK":
            sheet.cell(row_number, 8).fill = FAILURE_FILL
            sheet.cell(row_number, 9).fill = FAILURE_FILL
            sheet.cell(row_number, 10).fill = INFO_FILL
            if record["recovered_if_excluded"]:
                sheet.cell(row_number, 11).fill = RECOVERY_FILL
                sheet.cell(row_number, 11).font = Font(
                    color="375623",
                    bold=True,
                )

        for offset, parameter in enumerate(parameters, parameter_start):
            raw_value = record["params"].get(parameter)
            value = normalize_value(raw_value)
            cell = sheet.cell(row_number, offset, value)
            cell.border = TABLE_BORDER
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = "0.0000"
                cell.alignment = Alignment(horizontal="right")
            parameter_grade = grade_value(parameter, raw_value, cuts)
            if parameter_grade not in {"N/A", "MISSING"}:
                cell.fill = GRADE_FILL[parameter_grade]
                cell.font = Font(color=GRADE_FONT[parameter_grade])

    widths = {
        1: 18,
        2: 18,
        3: 20,
        4: 18,
        5: 20,
        6: 18,
        7: 16,
        8: 16,
        9: 28,
        10: 20,
        11: 42,
    }
    for column in range(
        parameter_start,
        parameter_start + len(parameters),
    ):
        widths[column] = 20
    configure_sheet(
        sheet,
        len(headers),
        "H5",
        widths,
        header_row=4,
    )
    sheet.sheet_properties.tabColor = "70AD47" if grade_key == "no_dac_grade" else "ED7D31"


def build_yield_stats(
    records: list[dict[str, Any]],
    grade_key: str,
) -> dict[str, Any]:
    counts = Counter(record[grade_key] for record in records)
    return {
        "total": len(records),
        "counts": {
            grade: counts.get(grade, 0)
            for grade in YIELD_GRADE_ORDER
        },
    }


def _write_yield_row(
    sheet,
    row: int,
    label: str,
    stats: dict[str, Any],
) -> None:
    start = 1
    sheet.cell(row, start, label)
    total_column = start + 1
    sheet.cell(row, total_column, stats["total"])
    sheet.cell(row, total_column).number_format = "#,##0"

    count_columns: list[int] = []
    for index, grade in enumerate(YIELD_GRADE_ORDER):
        count_column = start + 2 + (index * 2)
        percent_column = count_column + 1
        count_columns.append(count_column)
        count_cell = sheet.cell(
            row,
            count_column,
            stats["counts"].get(grade, 0),
        )
        count_cell.number_format = "#,##0"
        count_cell.fill = GRADE_FILL[grade]
        count_cell.font = Font(
            color=GRADE_FONT[grade],
            bold=True,
        )
        total_letter = get_column_letter(total_column)
        count_letter = get_column_letter(count_column)
        percent_cell = sheet.cell(
            row,
            percent_column,
            f'=IF(${total_letter}{row}=0,0,{count_letter}{row}/${total_letter}{row})',
        )
        percent_cell.number_format = "0.0%"
        percent_cell.fill = GRADE_FILL[grade]
        percent_cell.font = Font(color=GRADE_FONT[grade])

    usable_column = start + 14
    yield_column = start + 15
    usable_letters = ",".join(
        f"{get_column_letter(column)}{row}"
        for column in count_columns[: len(GRADE_ORDER)]
    )
    sheet.cell(row, usable_column, f"=SUM({usable_letters})")
    sheet.cell(row, usable_column).number_format = "#,##0"
    sheet.cell(row, usable_column).fill = PatternFill(
        "solid",
        fgColor="E2F0D9",
    )
    total_letter = get_column_letter(total_column)
    usable_letter = get_column_letter(usable_column)
    sheet.cell(
        row,
        yield_column,
        f'=IF(${total_letter}{row}=0,0,{usable_letter}{row}/${total_letter}{row})',
    )
    sheet.cell(row, yield_column).number_format = "0.0%"
    sheet.cell(row, yield_column).fill = PatternFill(
        "solid",
        fgColor="E2F0D9",
    )
    sheet.cell(row, usable_column).font = Font(
        color="375623",
        bold=True,
    )
    sheet.cell(row, yield_column).font = Font(
        color="375623",
        bold=True,
    )

    last_column = yield_column
    for column in range(1, last_column + 1):
        cell = sheet.cell(row, column)
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(
            horizontal="left" if column == start else "center",
            vertical="center",
        )


def write_yield_analysis(
    workbook,
    records: list[dict[str, Any]],
    batch_records: dict[str, list[dict[str, Any]]],
    criteria_version: int | None,
) -> None:
    if "Chip Yield Analysis" in workbook.sheetnames:
        del workbook["Chip Yield Analysis"]
    sheet = workbook.create_sheet("Chip Yield Analysis", 1)
    last_column = 16
    set_title(
        sheet,
        "ALTAI Chip Yield Analysis",
        (
            "Usable = GOLD + SILVER + BRONZE + BRONZENOVBB. "
            "CONTACT_FAIL is separated from NK. "
            f"Criteria version: {criteria_version if criteria_version is not None else 'unknown'}."
        ),
        last_column,
    )

    sheet.merge_cells("A4:P4")
    sheet["A4"] = "Current cumulative yield — all chips"
    sheet["A4"].fill = TITLE_FILL
    sheet["A4"].font = Font(color="FFFFFF", bold=True, size=12)
    current_headers = [
        "Test Mode",
        "Total",
        *(
            item
            for grade in YIELD_GRADE_ORDER
            for item in (f"{grade} Count", f"{grade} %")
        ),
        "Usable Count",
        "Yield %",
    ]
    for column, header in enumerate(current_headers, 1):
        style_header(sheet.cell(5, column, header))
    for row, (mode, grade_key) in enumerate(YIELD_MODES, 6):
        _write_yield_row(
            sheet,
            row,
            mode,
            build_yield_stats(records, grade_key),
        )

    sheet.merge_cells("A10:P10")
    sheet["A10"] = (
        "Yield by result sheet — one row per testdata batch, latest first"
    )
    sheet["A10"].fill = TITLE_FILL
    sheet["A10"].font = Font(color="FFFFFF", bold=True, size=12)
    batch_headers = [
        "Result Sheet",
        *current_headers[1:],
    ]
    for column, header in enumerate(batch_headers, 1):
        style_header(sheet.cell(11, column, header))

    row = 12
    for batch_name in sorted(batch_records, reverse=True):
        _write_yield_row(
            sheet,
            row,
            batch_result_sheet_name(batch_name),
            build_yield_stats(
                batch_records[batch_name],
                "no_dac_grade",
            ),
        )
        row += 1

    sheet.freeze_panes = "B12"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[5].height = 40
    sheet.row_dimensions[11].height = 40
    widths = {1: 34, 2: 15}
    for column in range(3, 17):
        widths[column] = 15
    widths[15] = 16
    widths[16] = 13
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.auto_filter.ref = f"A11:P{max(sheet.max_row, 11)}"
    sheet.sheet_properties.tabColor = "7030A0"


def build_failure_analysis_stats(
    records: list[dict[str, Any]],
    cuts: dict[str, dict[str, list[float]]],
) -> dict[str, Any]:
    total = len(records)
    current_nk = sum(
        record["no_dac_grade"] == "NK"
        for record in records
    )
    current_usable = sum(
        record["no_dac_grade"] in GRADE_ORDER
        for record in records
    )
    parameter_rows = []
    for parameter in no_dac_active_parameters(cuts):
        failing_records = [
            record
            for record in records
            if (
                record["no_dac_grade"] == "NK"
                and parameter in record["nk_failure_states"]
            )
        ]
        recovered_records = [
            record
            for record in records
            if (
                record["no_dac_grade"] == "NK"
                and parameter in record["recovery_by_parameter"]
            )
        ]
        recovered_grades = Counter(
            record["recovery_by_parameter"][parameter]
            for record in recovered_records
        )
        parameter_rows.append(
            {
                "parameter": parameter,
                "group": TEST_GROUP_LABELS.get(
                    parameter_group(parameter),
                    parameter_group(parameter),
                ),
                "failing": len(failing_records),
                "missing": sum(
                    record["nk_failure_states"].get(parameter) == "MISSING"
                    for record in failing_records
                ),
                "out_of_range": sum(
                    record["nk_failure_states"].get(parameter)
                    == "OUT_OF_RANGE"
                    for record in failing_records
                ),
                "recovered": len(recovered_records),
                "recovered_grades": {
                    grade: recovered_grades.get(grade, 0)
                    for grade in GRADE_ORDER
                },
            }
        )

    parameter_rows.sort(
        key=lambda item: (
            -item["recovered"],
            -item["failing"],
            TEST_GROUP_INDEX.get(
                parameter_group(item["parameter"]),
                999,
            ),
            item["parameter"],
        )
    )
    return {
        "total": total,
        "current_nk": current_nk,
        "current_usable": current_usable,
        "contact_fail": sum(
            record["status"] == "CONTACT_FAIL"
            for record in records
        ),
        "parameters": parameter_rows,
    }


def _write_failure_analysis_row(
    sheet,
    row: int,
    item: dict[str, Any],
    stats: dict[str, Any],
) -> None:
    values = [
        item["parameter"],
        item["group"],
        stats["total"],
        stats["current_nk"],
        item["failing"],
        None,
        item["missing"],
        item["out_of_range"],
        item["recovered"],
        item["recovered_grades"]["GOLD"],
        item["recovered_grades"]["SILVER"],
        item["recovered_grades"]["BRONZE"],
        item["recovered_grades"]["BRONZENOVBB"],
        stats["current_usable"],
        None,
        None,
        None,
        None,
    ]
    for column, value in enumerate(values, 1):
        cell = sheet.cell(row, column, value)
        cell.border = TABLE_BORDER
        cell.alignment = Alignment(
            horizontal="left" if column in {1, 2} else "center",
            vertical="center",
        )
        if column not in {1, 2, 6, 15, 16, 17, 18}:
            cell.number_format = "#,##0"

    sheet.cell(
        row,
        6,
        f'=IF($D{row}=0,0,$E{row}/$D{row})',
    )
    sheet.cell(
        row,
        15,
        f"=$N{row}+$I{row}",
    )
    sheet.cell(
        row,
        16,
        f'=IF($C{row}=0,0,$N{row}/$C{row})',
    )
    sheet.cell(
        row,
        17,
        f'=IF($C{row}=0,0,$O{row}/$C{row})',
    )
    sheet.cell(
        row,
        18,
        f"=$Q{row}-$P{row}",
    )
    for column in {6, 16, 17, 18}:
        sheet.cell(row, column).number_format = "0.0%"

    sheet.cell(row, 5).fill = FAILURE_FILL
    sheet.cell(row, 6).fill = FAILURE_FILL
    sheet.cell(row, 7).fill = INFO_FILL
    sheet.cell(row, 8).fill = FAILURE_FILL
    sheet.cell(row, 9).fill = RECOVERY_FILL
    sheet.cell(row, 10).fill = GRADE_FILL["GOLD"]
    sheet.cell(row, 11).fill = GRADE_FILL["SILVER"]
    sheet.cell(row, 12).fill = GRADE_FILL["BRONZE"]
    sheet.cell(row, 13).fill = GRADE_FILL["BRONZENOVBB"]
    for column in {14, 15, 16, 17, 18}:
        sheet.cell(row, column).fill = RECOVERY_FILL
    if item["recovered"] > 0:
        for column in {9, 15, 17, 18}:
            sheet.cell(row, column).font = Font(
                color="375623",
                bold=True,
            )


def write_failure_analysis_section(
    sheet,
    start_row: int,
    label: str,
    records: list[dict[str, Any]],
    cuts: dict[str, dict[str, list[float]]],
) -> int:
    stats = build_failure_analysis_stats(records, cuts)
    last_column = 18
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(
        start_row=start_row,
        start_column=1,
        end_row=start_row,
        end_column=last_column,
    )
    title_cell = sheet.cell(start_row, 1, label)
    title_cell.fill = TITLE_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=12)
    title_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[start_row].height = 24

    header_row = start_row + 1
    headers = [
        "Parameter",
        "Test Group",
        "Total Chips",
        "Current NK",
        "NK Chips Failing This Criterion",
        "% of Current NK",
        "Missing",
        "Out of Range",
        "Recovered if Excluded",
        "Recovered GOLD",
        "Recovered SILVER",
        "Recovered BRONZE",
        "Recovered BRONZENOVBB",
        "Current Usable",
        "Usable if Excluded",
        "Current Yield",
        "Yield if Excluded",
        "Yield Gain",
    ]
    for column, header in enumerate(headers, 1):
        style_header(sheet.cell(header_row, column, header))
    sheet.row_dimensions[header_row].height = 48

    row = header_row + 1
    for item in stats["parameters"]:
        _write_failure_analysis_row(
            sheet,
            row,
            item,
            stats,
        )
        row += 1

    note_row = row
    sheet.merge_cells(f"A{note_row}:{last_letter}{note_row}")
    sheet.cell(
        note_row,
        1,
        (
            f"Scope totals: {stats['total']} chips; "
            f"current NK {stats['current_nk']}; "
            f"current usable {stats['current_usable']}; "
            f"CONTACT_FAIL {stats['contact_fail']}. "
            "Yield uses all chips as the denominator, matching Chip Yield Analysis."
        ),
    )
    sheet.cell(note_row, 1).fill = SUBTITLE_FILL
    sheet.cell(note_row, 1).font = Font(
        color="1F1F1F",
        italic=True,
        size=9,
    )
    sheet.cell(note_row, 1).alignment = Alignment(
        vertical="center",
        wrap_text=True,
    )
    sheet.row_dimensions[note_row].height = 30
    return note_row


def write_failure_analysis(
    workbook,
    records: list[dict[str, Any]],
    batch_records: dict[str, list[dict[str, Any]]],
    cuts: dict[str, dict[str, list[float]]],
    criteria_version: int | None,
) -> None:
    if "Failure Analysis" in workbook.sheetnames:
        del workbook["Failure Analysis"]
    sheet = workbook.create_sheet("Failure Analysis", 2)
    last_column = 18
    set_title(
        sheet,
        "ALTAI no-DACSCAN Failure Analysis",
        (
            "Order-independent leave-one-out analysis: each active no-DACSCAN "
            "parameter is excluded alone and the chip is fully reclassified. "
            "DACScan and Backbias are not active causes; CONTACT_FAIL remains "
            "separate. Missing and out-of-range failures are shown separately. "
            f"Criteria version: {criteria_version if criteria_version is not None else 'unknown'}."
        ),
        last_column,
    )

    cumulative_end = write_failure_analysis_section(
        sheet,
        4,
        "Current cumulative failure analysis — all chips",
        records,
        cuts,
    )
    row = cumulative_end + 3
    for batch_name in sorted(batch_records, reverse=True):
        row = write_failure_analysis_section(
            sheet,
            row,
            (
                f"{batch_result_sheet_name(batch_name)} — "
                "failure analysis"
            ),
            batch_records[batch_name],
            cuts,
        ) + 3

    sheet.freeze_panes = "C6"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A5:R{cumulative_end - 1}"
    widths = {
        1: 44,
        2: 22,
        3: 13,
        4: 13,
        5: 20,
        6: 15,
        7: 13,
        8: 15,
        9: 18,
        10: 16,
        11: 16,
        12: 16,
        13: 19,
        14: 16,
        15: 18,
        16: 15,
        17: 18,
        18: 14,
    }
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.sheet_properties.tabColor = "C65911"


def write_nk_chip_details(
    workbook,
    batch_records: dict[str, list[dict[str, Any]]],
    cuts: dict[str, dict[str, list[float]]],
    criteria_version: int | None,
) -> None:
    if "NK Chip Details" in workbook.sheetnames:
        del workbook["NK Chip Details"]
    sheet = workbook.create_sheet("NK Chip Details", 3)
    headers = [
        "Data Batch",
        "Original Chip ID",
        "Wrong Written Chip ID",
        "Test Date",
        "Status",
        "Current Grade",
        "Parameter",
        "Test Group",
        "Failure Type",
        "Failed Grade Cuts",
        "Measured Value",
        "Recovered if Excluded",
        "Recovered Grade",
        "Source Folder",
    ]
    set_title(
        sheet,
        "ALTAI NK Chip Details — no-DACSCAN",
        (
            "One row per NK chip and failed active criterion. Use filters to "
            "trace missing/out-of-range measurements and leave-one-out recovery. "
            "DACScan and Backbias are excluded. CONTACT_FAIL is not listed as "
            f"an electrical NK cause. Criteria version: "
            f"{criteria_version if criteria_version is not None else 'unknown'}."
        ),
        len(headers),
    )
    write_group_header(sheet, 1, 6, "Chip & Test")
    write_group_header(sheet, 7, 11, "Failed Criterion")
    write_group_header(sheet, 12, 13, "Leave-One-Out Recovery")
    write_group_header(sheet, 14, 14, "Source")

    for column, header in enumerate(headers, 1):
        style_header(sheet.cell(4, column, header))

    row = 5
    for batch_name in sorted(batch_records, reverse=True):
        for record in batch_records[batch_name]:
            if record["no_dac_grade"] != "NK":
                continue
            if record["status"] != "ANALYZED":
                values = [
                    batch_name,
                    record["chip_id"],
                    record["original_chip_id"],
                    record["test_date"],
                    record["status"],
                    record["no_dac_grade"],
                    "(analysis output)",
                    "Test Metadata",
                    record["status"],
                    "",
                    None,
                    "NO",
                    "",
                    record["source"],
                ]
                detail_rows = [values]
            else:
                failed, states = grade_failure_details(
                    record["params"],
                    cuts,
                )
                detail_rows = []
                for parameter in no_dac_active_parameters(cuts):
                    if parameter not in failed:
                        continue
                    recovered_grade = record[
                        "recovery_by_parameter"
                    ].get(parameter, "")
                    detail_rows.append(
                        [
                            batch_name,
                            record["chip_id"],
                            record["original_chip_id"],
                            record["test_date"],
                            record["status"],
                            record["no_dac_grade"],
                            parameter,
                            TEST_GROUP_LABELS.get(
                                parameter_group(parameter),
                                parameter_group(parameter),
                            ),
                            states[parameter],
                            ", ".join(
                                GRADE_SHORT_LABELS[grade]
                                for grade in failed[parameter]["grades"]
                            ),
                            normalize_value(
                                record["params"].get(parameter)
                            ),
                            "YES" if recovered_grade else "NO",
                            recovered_grade,
                            record["source"],
                        ]
                    )

            for values in detail_rows:
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row, column, value)
                    cell.border = TABLE_BORDER
                    cell.alignment = Alignment(
                        horizontal=(
                            "left"
                            if column in {1, 2, 3, 7, 8, 14}
                            else "center"
                        ),
                        vertical="center",
                    )
                style_grade_cell(
                    sheet.cell(row, 6),
                    record["no_dac_grade"],
                )
                failure_type = sheet.cell(row, 9).value
                sheet.cell(row, 9).fill = (
                    INFO_FILL
                    if failure_type == "MISSING"
                    else FAILURE_FILL
                )
                measured_cell = sheet.cell(row, 11)
                if is_numeric_measurement(measured_cell.value):
                    measured_cell.number_format = "0.0000"
                if sheet.cell(row, 12).value == "YES":
                    sheet.cell(row, 12).fill = RECOVERY_FILL
                    sheet.cell(row, 12).font = Font(
                        color="375623",
                        bold=True,
                    )
                    style_grade_cell(
                        sheet.cell(row, 13),
                        sheet.cell(row, 13).value,
                    )
                else:
                    sheet.cell(row, 12).fill = INFO_FILL
                sheet.cell(row, 14).font = Font(
                    color="7F7F7F",
                    size=9,
                )
                row += 1

    configure_sheet(
        sheet,
        len(headers),
        "G5",
        {
            1: 20,
            2: 18,
            3: 20,
            4: 20,
            5: 18,
            6: 16,
            7: 46,
            8: 22,
            9: 18,
            10: 20,
            11: 18,
            12: 21,
            13: 18,
            14: 52,
        },
        header_row=4,
    )
    sheet.sheet_properties.tabColor = "A61C00"


def create_or_update_workbook(
    output_path: Path,
    records: list[dict[str, Any]],
    batch_records: dict[str, list[dict[str, Any]]],
    cuts: dict[str, dict[str, list[float]]],
    criteria_version: int | None,
    usage_source: Path | None,
) -> dict[str, dict[str, str]]:
    if output_path.is_file():
        workbook = load_workbook(output_path)
    else:
        workbook = Workbook()
        workbook.remove(workbook.active)

    existing_manual = extract_manual_fields(workbook)
    imported_manual = load_usage_source(usage_source)
    manual_fields = merge_manual_fields(imported_manual, existing_manual)

    for sheet_name in list(workbook.sheetnames):
        if (
            sheet_name.startswith("Only Power_")
            or sheet_name.startswith("no DACSCAN_")
        ):
            del workbook[sheet_name]

    write_all_chip_data(
        workbook,
        records,
        cuts,
        criteria_version,
        manual_fields,
    )
    write_yield_analysis(
        workbook,
        records,
        batch_records,
        criteria_version,
    )
    write_failure_analysis(
        workbook,
        records,
        batch_records,
        cuts,
        criteria_version,
    )
    write_nk_chip_details(
        workbook,
        batch_records,
        cuts,
        criteria_version,
    )

    parameters = all_parameters(cuts)
    for index, batch_name in enumerate(
        sorted(batch_records, reverse=True),
        start=4,
    ):
        write_snapshot(
            workbook,
            batch_result_sheet_name(batch_name),
            index,
            (
                f"ALTAI Final Results — {batch_name} — "
                f"DACScan & Backbias Excluded "
                f"(PowerTest.i_d −4.4 mA)"
            ),
            batch_records[batch_name],
            "no_dac_grade",
            parameters,
            cuts,
            criteria_version,
            excluded_groups=NO_DAC_EXCLUDED_GROUPS,
        )

    RESULTS_ROOT.mkdir(parents=True, exist_ok=True)
    temporary = output_path.with_name(f".{output_path.stem}.tmp.xlsx")
    try:
        workbook.save(temporary)
        os.replace(temporary, output_path)
    finally:
        if temporary.exists():
            temporary.unlink()
    return manual_fields


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Create/update the cumulative ALTAI chip-test workbook using only "
            "ALTAI_CHIPTEST_NEW/testdata."
        )
    )
    parser.add_argument(
        "-o",
        "--output-name",
        default=DEFAULT_OUTPUT_NAME,
        help=(
            "Output .xlsx file name inside Results "
            f"(default: {DEFAULT_OUTPUT_NAME})"
        ),
    )
    parser.add_argument(
        "--usage-source",
        type=Path,
        help=(
            "Optional existing usage-record .xlsx file. Imported values seed "
            "Usage and Current Location; edits already present in the output "
            "workbook take precedence."
        ),
    )
    args = parser.parse_args()

    validate_project_layout()
    output_path = get_output_path(args.output_name)
    cuts, criteria_version = load_criteria()

    all_records = collect_records()
    if not all_records:
        raise RuntimeError(f"No test-run directories found below {TESTDATA_ROOT}")
    final_records = add_grades(latest_per_chip(all_records), cuts)
    batch_records = build_batch_records(all_records, cuts)

    manual_fields = create_or_update_workbook(
        output_path,
        final_records,
        batch_records,
        cuts,
        criteria_version,
        args.usage_source,
    )

    status_counts = Counter(record["status"] for record in final_records)
    no_dac_counts = Counter(record["no_dac_grade"] for record in final_records)
    power_counts = Counter(record["power_grade"] for record in final_records)

    print(f"Project testdata: {TESTDATA_ROOT}")
    print(f"Test batches: {len(list(TESTDATA_ROOT.glob('testdata_*')))}")
    print(f"All test runs: {len(all_records)}")
    print(f"Final unique chips: {len(final_records)}")
    print(f"Status: {dict(status_counts)}")
    print(f"no DACSCAN: {dict(no_dac_counts)}")
    print(f"Only Power: {dict(power_counts)}")
    for batch_name in sorted(batch_records, reverse=True):
        batch_status = Counter(
            record["status"]
            for record in batch_records[batch_name]
        )
        batch_no_dac = Counter(
            record["no_dac_grade"]
            for record in batch_records[batch_name]
        )
        print(
            f"{batch_name}: chips={len(batch_records[batch_name])}, "
            f"status={dict(batch_status)}, "
            f"no DACSCAN={dict(batch_no_dac)}"
        )
    print(
        "Usage/location records: "
        f"{sum(bool(v.get('usage') or v.get('location')) for v in manual_fields.values())}"
    )
    print(f"Criteria version: {criteria_version}")
    print(f"Workbook: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
